# -*- coding: utf-8 -*-
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import List, Optional, Dict
from passlib.context import CryptContext
import os
import io
import json
import time
import logging
import requests
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from urllib.parse import quote
from dotenv import load_dotenv

import psycopg2
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger(__name__)

load_dotenv()

GOOGLE_MAPS_API_KEY = os.environ.get("GOOGLE_MAPS_API_KEY", "")
FIXSTARS_API_KEY = os.environ.get("FIXSTARS_API_KEY", "")
NAVITIME_API_KEY = os.environ.get("NAVITIME_API_KEY", "")

logger.info(f"FIXSTARS_API_KEY: {'設定あり' if FIXSTARS_API_KEY else '未設定'}")
logger.info(f"NAVITIME_API_KEY: {'設定あり' if NAVITIME_API_KEY else '未設定'}")
logger.info(f"GOOGLE_MAPS_API_KEY: {'設定あり' if GOOGLE_MAPS_API_KEY else '未設定'}")

# ==========================================
# PostgreSQL 接続
# ==========================================
POSTGRES_HOST = os.environ.get("POSTGRES_HOST", "db")
POSTGRES_PORT = int(os.environ.get("POSTGRES_PORT", "5432"))
POSTGRES_USER = os.environ.get("POSTGRES_USER", "user")
POSTGRES_PASSWORD = os.environ.get("POSTGRES_PASSWORD", "password")
POSTGRES_DB = os.environ.get("POSTGRES_DB", "circledb")

def get_db_connection():
    return psycopg2.connect(
        host=POSTGRES_HOST, port=POSTGRES_PORT,
        user=POSTGRES_USER, password=POSTGRES_PASSWORD, dbname=POSTGRES_DB
    )

def init_db():
    try:
        conn = get_db_connection()
        with conn.cursor() as cur:
            cur.execute("""
                CREATE TABLE IF NOT EXISTS distance_cache (
                    station_a VARCHAR(200) NOT NULL,
                    station_b VARCHAR(200) NOT NULL,
                    time_slot VARCHAR(30) NOT NULL,
                    duration_minutes INTEGER NOT NULL,
                    PRIMARY KEY (station_a, station_b, time_slot)
                )
            """)
            cur.execute("""
                CREATE TABLE IF NOT EXISTS coord_cache (
                    station_name VARCHAR(200) NOT NULL PRIMARY KEY,
                    lat DOUBLE PRECISION NOT NULL,
                    lon DOUBLE PRECISION NOT NULL
                )
            """)
            cur.execute("""
                CREATE TABLE IF NOT EXISTS users (
                    id SERIAL PRIMARY KEY,
                    email VARCHAR(255) UNIQUE NOT NULL,
                    password_hash VARCHAR(255) NOT NULL,
                    name VARCHAR(255) NOT NULL,
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                )
            """)
        conn.commit()
        conn.close()
        logger.info("DBテーブルの初期化完了")
    except Exception as e:
        logger.warning(f"DB初期化失敗（JSONフォールバックで継続）: {e}")

def _load_distance_cache_from_db() -> dict:
    try:
        conn = get_db_connection()
        cache = {}
        with conn.cursor() as cur:
            cur.execute("SELECT station_a, station_b, time_slot, duration_minutes FROM distance_cache")
            for row in cur.fetchall():
                cache[f"{row[0]}_{row[1]}_{row[2]}"] = row[3]
        conn.close()
        logger.info(f"distance_cache: {len(cache)}件をDBから読み込み")
        return cache
    except Exception as e:
        logger.warning(f"distance_cacheのDB読み込み失敗: {e}")
        return {}

def _load_coord_cache_from_db() -> dict:
    try:
        conn = get_db_connection()
        cache = {}
        with conn.cursor() as cur:
            cur.execute("SELECT station_name, lat, lon FROM coord_cache")
            for row in cur.fetchall():
                cache[row[0]] = {"lat": row[1], "lon": row[2]}
        conn.close()
        logger.info(f"coord_cache: {len(cache)}件をDBから読み込み")
        return cache
    except Exception as e:
        logger.warning(f"coord_cacheのDB読み込み失敗: {e}")
        return {}

def _save_distance_to_db(station_a: str, station_b: str, time_slot: str, duration: int):
    try:
        conn = get_db_connection()
        with conn.cursor() as cur:
            cur.execute("""
                INSERT INTO distance_cache (station_a, station_b, time_slot, duration_minutes)
                VALUES (%s, %s, %s, %s)
                ON CONFLICT (station_a, station_b, time_slot) DO UPDATE
                SET duration_minutes = EXCLUDED.duration_minutes
            """, (station_a, station_b, time_slot, duration))
        conn.commit()
        conn.close()
    except Exception as e:
        logger.warning(f"distance_cacheのDB保存失敗: {e}")

def _save_coord_to_db(station_name: str, lat: float, lon: float):
    try:
        conn = get_db_connection()
        with conn.cursor() as cur:
            cur.execute("""
                INSERT INTO coord_cache (station_name, lat, lon)
                VALUES (%s, %s, %s)
                ON CONFLICT (station_name) DO UPDATE
                SET lat = EXCLUDED.lat, lon = EXCLUDED.lon
            """, (station_name, lat, lon))
        conn.commit()
        conn.close()
    except Exception as e:
        logger.warning(f"coord_cacheのDB保存失敗: {e}")
NAVITIME_API_HOST = "navitime-route-totalnavi.p.rapidapi.com"

app = FastAPI()

@app.on_event("startup")
def startup_event():
    global _distance_cache, _coord_cache
    init_db()
    _distance_cache = _load_distance_cache_from_db()
    _coord_cache = _load_coord_cache_from_db()

ALLOWED_ORIGINS = os.environ.get(
    "ALLOWED_ORIGINS",
    "https://circlemanage.vercel.app,http://localhost:3000"
).split(",")

app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS,
    # ブランチごとに変わるVercelプレビューURL(circlemanage-git-xxx.vercel.app等)を
    # 都度ALLOWED_ORIGINSに追加しなくて済むよう、vercel.appサブドメインは包括的に許可する
    allow_origin_regex=r"https://.*\.vercel\.app",
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ==========================================
# 入力データの構造定義
# ==========================================
class Member(BaseModel):
    name: str
    station: str
    can_drive: bool
    capacity: Optional[int] = None
    want_with: Optional[List[str]] = []
    awkward_with: Optional[List[str]] = []

class EventData(BaseModel):
    members: List[Member]
    target_arrival: Optional[str] = ""
    p_score: Optional[int] = -5

class FormConfig(BaseModel):
    access_token: str
    event_name: str = "イベント参加フォーム"

class SheetConfig(BaseModel):
    access_token: str
    spreadsheet_id: str

class CampRosterMember(BaseModel):
    grade: str = ""
    name: str

class CampCostItem(BaseModel):
    label: str
    amount: float = 0

class CampSpecialDinnerPrice(BaseModel):
    date: str  # ISO形式 (yyyy-mm-dd)。BBQ等でこの日だけ夕食が特別単価になる
    price: float = 0

class CampRosterExportRequest(BaseModel):
    members: List[CampRosterMember]
    dates: List[str]  # ISO形式 (yyyy-mm-dd)
    # 出欠: 氏名 -> 日付 -> 食事("朝"/"昼"/"夜") -> 参加有無
    attendance: Dict[str, Dict[str, Dict[str, bool]]] = {}
    lodging_fee: float = 0  # 宿泊費(1人1泊あたり)
    meal_prices: Dict[str, float] = {}  # "朝"/"昼"/"夜" -> 1食あたり単価
    special_dinner_prices: List[CampSpecialDinnerPrice] = []  # BBQ等、特定日だけ夕食単価が異なる場合
    deposit_amount: float = 0  # 前金(1人あたり、先に集める金額)
    deposit_paid: Dict[str, bool] = {}  # 氏名 -> 前金を徴収済みか
    cost_items: List[CampCostItem] = []

class CostMember(BaseModel):
    name: str
    can_drive: bool
    has_insurance: bool = False
    pre_paid: bool = False
    advance_payment: float = 0

class UserRegister(BaseModel):
    email: str
    password: str
    name: str

class UserVerify(BaseModel):
    email: str
    password: str

class CostConfig(BaseModel):
    members: List[CostMember]
    participate: int = 11500
    pre_payed: int = 15000
    entry: int = 20000
    alcohol: List[int] = []
    petlorem: List[int] = []
    express: List[int] = []
    lent: List[int] = []
    finance: int = 600

# ==========================================
# ユーザー認証エンドポイント
# ==========================================
@app.post("/auth/register")
async def register_user(data: UserRegister):
    if len(data.password.encode("utf-8")) > 72:
        raise HTTPException(status_code=400, detail="パスワードは72文字以内にしてください")
    try:
        conn = get_db_connection()
        with conn.cursor() as cur:
            cur.execute("SELECT id FROM users WHERE email = %s", (data.email,))
            if cur.fetchone():
                raise HTTPException(status_code=400, detail="このメールアドレスは既に登録されています")
            password_hash = pwd_context.hash(data.password)
            cur.execute(
                "INSERT INTO users (email, password_hash, name) VALUES (%s, %s, %s) RETURNING id",
                (data.email, password_hash, data.name)
            )
            user_id = cur.fetchone()[0]
        conn.commit()
        conn.close()
        return {"id": str(user_id), "email": data.email, "name": data.name}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"ユーザー登録エラー: {e}")
        raise HTTPException(status_code=500, detail="登録に失敗しました")

@app.post("/auth/verify")
async def verify_user(data: UserVerify):
    try:
        conn = get_db_connection()
        with conn.cursor() as cur:
            cur.execute("SELECT id, password_hash, name FROM users WHERE email = %s", (data.email,))
            row = cur.fetchone()
        conn.close()
        if not row or not pwd_context.verify(data.password, row[1]):
            raise HTTPException(status_code=401, detail="メールアドレスまたはパスワードが違います")
        return {"id": str(row[0]), "email": data.email, "name": row[2]}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"認証エラー: {e}")
        raise HTTPException(status_code=500, detail="認証に失敗しました")


# ==========================================
# 費用計算エンドポイント
# ==========================================
@app.post("/calculate-costs")
async def calculate_costs(data: CostConfig):
    drivers = [m for m in data.members if m.can_drive]
    passengers = [m for m in data.members if not m.can_drive]

    if not drivers:
        return {"error": "運転手が1人もいません。"}
    if not passengers:
        return {"error": "乗客が1人もいません。"}

    sum_petlorem = sum(data.petlorem)
    sum_express = sum(data.express)
    sum_lent = sum(data.lent)
    sum_alcohol = sum(data.alcohol)
    total = len(drivers) + len(passengers)

    pass_val = (
        (sum_petlorem + sum_express + sum_lent) / len(passengers)
        + (sum_alcohol + data.entry) / total
        + data.participate
    )
    drive_val = (sum_alcohol + data.entry) / total + data.participate

    costs = []
    for m in data.members:
        pre_check = data.pre_payed if m.pre_paid else 0
        base = drive_val if m.can_drive else pass_val
        if m.has_insurance:
            amount = int(data.finance + base) + 1
        else:
            amount = int(base) + 1
        amount -= pre_check
        amount -= int(m.advance_payment)
        costs.append({
            "name": m.name,
            "role": "運転手" if m.can_drive else "乗客",
            "amount": amount,
        })

    return {
        "costs": costs,
        "summary": {
            "drivers": len(drivers),
            "passengers": len(passengers),
            "pass_val": round(pass_val, 1),
            "drive_val": round(drive_val, 1),
        },
    }


# ==========================================
# 配車計算エンドポイント
# ==========================================
@app.post("/assign")
async def assign_members(data: EventData):
    members = data.members
    drivers = [m for m in members if m.can_drive]
    passengers = [m for m in members if not m.can_drive]

    if not drivers:
        return {"error": "運転手が1人もいません。"}
    if not passengers:
        return {"error": "乗客が1人もいません。"}

    C = [(d.capacity or 4) for d in drivers]
    total_seats = sum(c - 1 for c in C)

    if len(passengers) > total_seats:
        return {"error": f"有効シート数({total_seats}席)が乗客数({len(passengers)}名)より少ないため、全員を配車できません。"}

    for m in members:
        m.station = normalize_station(m.station)

    W, driver_W = build_relation_matrix(passengers, drivers, data.p_score or -5)

    d_matrix = None
    navitime_used = False
    navitime_error = None
    if not NAVITIME_API_KEY:
        navitime_error = "NAVITIME_API_KEY が未設定"
    elif not GOOGLE_MAPS_API_KEY:
        navitime_error = "GOOGLE_MAPS_API_KEY が未設定"
    else:
        try:
            target_arrival = None
            if data.target_arrival:
                try:
                    target_arrival = datetime.fromisoformat(data.target_arrival)
                except ValueError:
                    pass
            if target_arrival is None:
                target_arrival = datetime.now().replace(hour=8, minute=0, second=0, microsecond=0)
            d_matrix = get_distance_matrix_navitime(
                [p.station for p in passengers],
                [d.station for d in drivers],
                target_arrival
            )
            navitime_used = True
        except Exception as e:
            d_matrix = None
            navitime_error = str(e)

    if d_matrix is None:
        d_matrix = [[0] * len(drivers) for _ in range(len(passengers))]

    distance_method = "navitime" if navitime_used else "greedy"

    amplify_error = None
    if not FIXSTARS_API_KEY:
        amplify_error = "FIXSTARS_API_KEY が未設定"
        logger.warning("FIXSTARS_API_KEY が未設定のためグリーディーにフォールバック")
    else:
        logger.info("Amplify 最適化を開始します")
        try:
            result = run_optimization_amplify(
                passengers, drivers, d_matrix, C, W, driver_W, FIXSTARS_API_KEY
            )
            logger.info(f"Amplify 最適化成功: feasible={result.get('feasible')}, objective={result.get('objective')}")
            result["distance_method"] = distance_method
            result["distance_error"] = navitime_error
            result["amplify_error"] = None
            return result
        except Exception as e:
            amplify_error = str(e)
            logger.error(f"Amplify 最適化失敗: {e}", exc_info=True)

    logger.info("グリーディー配車を実行します")
    result = run_greedy_assignment(passengers, drivers, d_matrix, C, W, driver_W)
    result["distance_method"] = distance_method
    result["distance_error"] = navitime_error
    result["amplify_error"] = amplify_error
    return result


# ==========================================
# 人間関係行列の生成
# ==========================================
def build_relation_matrix(
    passengers: List[Member],
    drivers: List[Member],
    p_score: int
):
    n_passengers = len(passengers)
    n_drivers = len(drivers)

    p_idx_map = {p.name: i for i, p in enumerate(passengers)}
    d_idx_map = {d.name: k for k, d in enumerate(drivers)}

    pref_p2p = {i: {j: 'NEUTRAL' for j in range(n_passengers)} for i in range(n_passengers)}
    pref_p2d = {i: {k: 'NEUTRAL' for k in range(n_drivers)} for i in range(n_passengers)}
    pref_d2p = {k: {i: 'NEUTRAL' for i in range(n_passengers)} for k in range(n_drivers)}

    for i, p in enumerate(passengers):
        for person in (p.want_with or []):
            person = person.strip()
            if person in p_idx_map:
                pref_p2p[i][p_idx_map[person]] = 'WANT'
            elif person in d_idx_map:
                pref_p2d[i][d_idx_map[person]] = 'WANT'
        for person in (p.awkward_with or []):
            person = person.strip()
            if person in p_idx_map:
                pref_p2p[i][p_idx_map[person]] = 'AWKWARD'
            elif person in d_idx_map:
                pref_p2d[i][d_idx_map[person]] = 'AWKWARD'

    for k, d in enumerate(drivers):
        for person in (d.want_with or []):
            person = person.strip()
            if person in p_idx_map:
                pref_d2p[k][p_idx_map[person]] = 'WANT'
        for person in (d.awkward_with or []):
            person = person.strip()
            if person in p_idx_map:
                pref_d2p[k][p_idx_map[person]] = 'AWKWARD'

    W = [[0] * n_passengers for _ in range(n_passengers)]
    for i in range(n_passengers):
        for j in range(i + 1, n_passengers):
            p1 = pref_p2p[i][j]
            p2 = pref_p2p[j][i]
            if p1 == 'AWKWARD' or p2 == 'AWKWARD':
                score = 100
            elif p1 == 'WANT' or p2 == 'WANT':
                score = p_score
            else:
                score = 0
            W[i][j] = score
            W[j][i] = score

    driver_W = [[0] * n_drivers for _ in range(n_passengers)]
    for i in range(n_passengers):
        for k in range(n_drivers):
            p1 = pref_p2d[i][k]
            p2 = pref_d2p[k][i]
            if p1 == 'AWKWARD' or p2 == 'AWKWARD':
                score = 100
            elif p1 == 'WANT' or p2 == 'WANT':
                score = p_score
            else:
                score = 0
            driver_W[i][k] = score

    return W, driver_W


# ==========================================
# NAVITIME API による所要時間行列取得（キャッシュ付き）
# ==========================================
USAGE_FILE = 'usage_stats.json'
RATE_LIMIT_PER_MIN = 50

def load_json(filename, default):
    if os.path.exists(filename):
        with open(filename, 'r', encoding='utf-8') as f:
            return json.load(f)
    return default

def save_json(filename, data):
    with open(filename, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=4)

# startup イベントで初期化される（モジュール読み込み時点では空）
_distance_cache: dict = {}
_coord_cache: dict = {}
_usage_stats: dict = load_json(USAGE_FILE, {"navitime_calls": 0})

def _round_to_30min(dt: datetime) -> str:
    rounded = (dt.minute // 30) * 30
    return dt.strftime(f"%Y-%m-%d %H:") + f"{rounded:02d}"

def _prefetch_coords(stations: list) -> None:
    uncached = [s for s in stations if s not in _coord_cache]
    if not uncached:
        return

    import googlemaps
    gmaps = googlemaps.Client(key=GOOGLE_MAPS_API_KEY)

    def fetch_one(name):
        result = gmaps.geocode(name + "駅")
        if not result:
            return name, None
        loc = result[0]['geometry']['location']
        return name, {"lat": loc['lat'], "lon": loc['lng']}

    with ThreadPoolExecutor(max_workers=5) as executor:
        for name, coord in executor.map(fetch_one, uncached):
            if coord:
                _coord_cache[name] = coord
                _save_coord_to_db(name, coord["lat"], coord["lon"])

def normalize_station(name: str) -> str:
    return name.strip().rstrip("駅")

def get_coords_with_cache(station_name: str) -> Optional[dict]:
    if station_name in _coord_cache:
        return _coord_cache[station_name]

    import googlemaps
    gmaps = googlemaps.Client(key=GOOGLE_MAPS_API_KEY)
    result = gmaps.geocode(station_name + "駅")
    if not result:
        return None

    loc = result[0]['geometry']['location']
    coord = {"lat": loc['lat'], "lon": loc['lng']}
    _coord_cache[station_name] = coord
    _save_coord_to_db(station_name, coord["lat"], coord["lon"])
    return coord

def call_navitime_api(start_coord: dict, goal_coord: dict, target_arrival: datetime) -> Optional[int]:
    url = f"https://{NAVITIME_API_HOST}/route_transit"
    arrival_str = target_arrival.strftime('%Y-%m-%dT%H:%M:%S')

    response = requests.get(
        url,
        headers={"X-RapidAPI-Key": NAVITIME_API_KEY, "X-RapidAPI-Host": NAVITIME_API_HOST},
        params={
            "start": f"{start_coord['lat']},{start_coord['lon']}",
            "goal": f"{goal_coord['lat']},{goal_coord['lon']}",
            "goal_time": arrival_str,
            "limit": "1"
        }
    )

    if response.status_code != 200:
        return None

    try:
        data = response.json()
        sections = data['items'][0].get('sections', [])
        train_sections = [
            s for s in sections
            if s.get('type') == 'move'
            and s.get('move') != 'walk'
            and s.get('line_name') != '徒歩'
        ]
        if not train_sections:
            return data['items'][0]['summary']['move']['time']

        start_dt = datetime.fromisoformat(train_sections[0]['from_time'].replace('Z', '+00:00'))
        end_dt = datetime.fromisoformat(train_sections[-1]['to_time'].replace('Z', '+00:00'))
        return int((end_dt - start_dt).total_seconds() / 60)
    except (KeyError, IndexError, ValueError):
        return None

def get_distance_matrix_navitime(
    passenger_stations: List[str],
    driver_stations: List[str],
    target_arrival: datetime
) -> List[List[int]]:
    time_slot = _round_to_30min(target_arrival)

    # ④ 未キャッシュ駅を並列ジオコーディング
    _prefetch_coords(list(set(passenger_stations + driver_stations)))

    # ③ ユニークなペアのみAPIを呼び出す（①時刻スロット付きキー、②対称ルート流用）
    unique_pairs = []
    for p in passenger_stations:
        for d in driver_stations:
            if p == d:
                continue
            if f"{p}_{d}_{time_slot}" not in _distance_cache and f"{d}_{p}_{time_slot}" not in _distance_cache:
                if (p, d) not in unique_pairs and (d, p) not in unique_pairs:
                    unique_pairs.append((p, d))

    api_counter = 0
    for p_station, d_station in unique_pairs:
        if api_counter > 0 and api_counter % RATE_LIMIT_PER_MIN == 0:
            time.sleep(60)

        p_coord = _coord_cache.get(p_station)
        d_coord = _coord_cache.get(d_station)

        if not p_coord or not d_coord:
            duration = 999
        else:
            duration = call_navitime_api(p_coord, d_coord, target_arrival)
            if duration is None:
                duration = 999
            api_counter += 1
            _usage_stats["navitime_calls"] += 1

        _distance_cache[f"{p_station}_{d_station}_{time_slot}"] = duration
        _save_distance_to_db(p_station, d_station, time_slot, duration)

    if unique_pairs:
        save_json(USAGE_FILE, _usage_stats)

    # マトリックス構築（②対称ルート流用）
    d_matrix = []
    for p_station in passenger_stations:
        row = []
        for d_station in driver_stations:
            if p_station == d_station:
                row.append(0)
                continue
            fwd = f"{p_station}_{d_station}_{time_slot}"
            rev = f"{d_station}_{p_station}_{time_slot}"
            row.append(_distance_cache.get(fwd) or _distance_cache.get(rev) or 999)
        d_matrix.append(row)

    return d_matrix


# ==========================================
# Amplify による最適化配車
# ==========================================
def run_optimization_amplify(
    passengers: List[Member],
    drivers: List[Member],
    d_matrix: List[List[int]],
    C: List[int],
    W: List[List[int]],
    driver_W: List[List[int]],
    fixstars_key: str
) -> dict:
    from amplify import VariableGenerator, Model, solve, equal_to, less_equal
    from amplify.client import FixstarsClient
    from amplify import sum as asum

    num_people = len(passengers)
    num_cars = len(drivers)

    alpha = 1.0
    beta = 1.0
    lambda_1 = 5000.0
    lambda_2 = 5000.0

    gen = VariableGenerator()
    x = gen.array("Binary", (num_people, num_cars))

    distance_cost = asum(
        d_matrix[i][k] * x[i, k]
        for i in range(num_people)
        for k in range(num_cars)
    )
    relation_cost = asum(
        W[i][j] * x[i, k] * x[j, k]
        for i in range(num_people)
        for j in range(i + 1, num_people)
        for k in range(num_cars)
    )
    driver_relation_cost = asum(
        driver_W[i][k] * x[i, k]
        for i in range(num_people)
        for k in range(num_cars)
    )
    objective = alpha * distance_cost + beta * relation_cost + beta * driver_relation_cost

    one_hot_constraints = [
        equal_to(asum(x[i, k] for k in range(num_cars)), 1)
        for i in range(num_people)
    ]
    capacity_constraints = [
        less_equal(asum(x[i, k] for i in range(num_people)), C[k] - 1)
        for k in range(num_cars)
    ]
    constraints = asum(c * lambda_1 for c in one_hot_constraints) + \
                  asum(c * lambda_2 for c in capacity_constraints)

    model = Model(objective, constraints)
    client = FixstarsClient()
    client.token = fixstars_key
    client.parameters.timeout = 1500

    result = solve(model, client)
    if len(result) == 0:
        return {"error": "Amplifyで解が見つかりませんでした。"}

    best = result[0]
    x_opt = x.evaluate(best.values)

    assignments = []
    assigned_ids = set()
    for k in range(num_cars):
        car_passengers = []
        for i in range(num_people):
            if int(x_opt[i][k]) == 1:
                car_passengers.append(f"{passengers[i].name} ({passengers[i].station}駅)")
                assigned_ids.add(i)
        assignments.append({
            "car_id": k + 1,
            "driver": f"{drivers[k].name} ({drivers[k].station}駅)",
            "members": [f"{drivers[k].name} (運転手 - {drivers[k].station}駅)"] + car_passengers,
            "feasible": best.feasible
        })

    unassigned = [passengers[i].name for i in range(num_people) if i not in assigned_ids]
    return {
        "assignments": assignments,
        "unassigned": unassigned,
        "method": "amplify",
        "objective": float(best.objective),
        "feasible": best.feasible
    }


# ==========================================
# フォールバック：グリーディ配車
# ==========================================
def run_greedy_assignment(
    passengers: List[Member],
    drivers: List[Member],
    d_matrix: List[List[int]],
    C: List[int],
    W: List[List[int]],
    driver_W: List[List[int]]
) -> dict:
    num_people = len(passengers)
    num_cars = len(drivers)

    assignment = [-1] * num_people
    car_counts = [0] * num_cars

    for i in range(num_people):
        best_car = -1
        best_score = float('inf')
        for k in range(num_cars):
            if car_counts[k] >= C[k] - 1:
                continue
            score = d_matrix[i][k] + driver_W[i][k]
            for j in range(num_people):
                if assignment[j] == k:
                    score += W[i][j]
            if score < best_score:
                best_score = score
                best_car = k
        if best_car != -1:
            assignment[i] = best_car
            car_counts[best_car] += 1

    assignments = []
    for k in range(num_cars):
        car_passengers = [
            f"{passengers[i].name} ({passengers[i].station}駅)"
            for i in range(num_people) if assignment[i] == k
        ]
        assignments.append({
            "car_id": k + 1,
            "driver": f"{drivers[k].name} ({drivers[k].station}駅)",
            "members": [f"{drivers[k].name} (運転手 - {drivers[k].station}駅)"] + car_passengers
        })

    unassigned = [passengers[i].name for i in range(num_people) if assignment[i] == -1]
    return {
        "assignments": assignments,
        "unassigned": unassigned,
        "method": "greedy"
    }


# ==========================================
# Google Forms API - フォーム自動作成
# ==========================================
@app.post("/create-form")
async def create_form(config: FormConfig):
    import requests

    headers = {
        "Authorization": f"Bearer {config.access_token}",
        "Content-Type": "application/json"
    }

    form_body = {
        "info": {
            "title": config.event_name,
            "documentTitle": config.event_name
        }
    }

    res = requests.post(
        "https://forms.googleapis.com/v1/forms",
        headers=headers,
        json=form_body
    )

    if res.status_code != 200:
        return {"error": f"フォーム作成に失敗しました: {res.text}"}

    form = res.json()
    form_id = form["formId"]

    questions = [
        {"title": "名前", "required": True, "type": "SHORT_ANSWER"},
        {"title": "最寄り駅", "required": True, "type": "SHORT_ANSWER"},
        {"title": "参加形態", "required": True, "type": "RADIO", "options": ["運転手", "乗客"]},
        {"title": "定員（運転手の方のみ・数字で入力）", "required": False, "type": "SHORT_ANSWER"},
        {"title": "一緒になりたい人（カンマ区切り）", "required": False, "type": "SHORT_ANSWER"},
        {"title": "気まずい人（カンマ区切り）", "required": False, "type": "SHORT_ANSWER"},
    ]

    requests_body = {"requests": []}
    for idx, q in enumerate(questions):
        item = {
            "createItem": {
                "item": {
                    "title": q["title"],
                    "questionItem": {
                        "question": {"required": q["required"]}
                    }
                },
                "location": {"index": idx}
            }
        }
        if q["type"] == "SHORT_ANSWER":
            item["createItem"]["item"]["questionItem"]["question"]["textQuestion"] = {}
        elif q["type"] == "RADIO":
            item["createItem"]["item"]["questionItem"]["question"]["choiceQuestion"] = {
                "type": "RADIO",
                "options": [{"value": opt} for opt in q["options"]]
            }
        requests_body["requests"].append(item)

    batch_res = requests.post(
        f"https://forms.googleapis.com/v1/forms/{form_id}:batchUpdate",
        headers=headers,
        json=requests_body
    )

    if batch_res.status_code != 200:
        return {"error": f"質問の追加に失敗しました: {batch_res.text}"}

    return {
        "form_id": form_id,
        "form_url": f"https://docs.google.com/forms/d/{form_id}/viewform",
        "edit_url": f"https://docs.google.com/forms/d/{form_id}/edit",
        "sheet_url": f"https://docs.google.com/forms/d/{form_id}/edit#responses",
    }


# ==========================================
# 合宿 参加者名簿 Excel出力
# 「2024合宿会計夏最終.xlsx」のSheet1（名簿）・合宿全体費用タブを
# ひな形として、学年/氏名/日付ごとの朝昼夜の出欠・泊数・徴収金額を
# 関数付きで出力する。
# ==========================================
CAMP_MEALS = ["朝", "昼", "夜"]

def _camp_date_label(iso: str) -> str:
    try:
        dt = datetime.strptime(iso, "%Y-%m-%d")
        return f"{dt.month}/{dt.day}"
    except Exception:
        return iso

@app.post("/export-camp-roster")
async def export_camp_roster(data: CampRosterExportRequest):
    wb = openpyxl.Workbook()

    # ---- スタイル定義 ----
    input_fill = PatternFill("solid", fgColor="DDEBF7")   # 青地: 自由に入力できる項目
    header_fill = PatternFill("solid", fgColor="404040")
    header_font = Font(color="FFFFFF", bold=True)
    calc_font = Font(color="217346", bold=True)            # 緑字: 数式で自動計算される項目
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    n_members = len(data.members)
    roster_start_row = 4  # 名簿シートのデータ開始行(ヘッダー3行: 日付/食事名/単価 の下)
    roster_last_row = roster_start_row + max(n_members, 1) - 1

    # ==================================================
    # シート1: 合宿全体費用（費用パラメータ）
    # ==================================================
    ws_cost = wb.active
    ws_cost.title = "合宿全体費用"

    ws_cost["A1"] = "色分け説明"
    ws_cost["A1"].font = Font(bold=True)
    ws_cost["A2"] = "青地：自由に金額を入力できる項目"
    ws_cost["A2"].fill = input_fill
    ws_cost["A3"] = "緑字：数式で自動計算される項目（編集しない）"
    ws_cost["A3"].font = calc_font

    ws_cost["A5"] = "参加人数"
    ws_cost["B5"] = f"=COUNTA(名簿!B{roster_start_row}:B{roster_last_row})"
    ws_cost["B5"].font = calc_font

    ws_cost["A6"] = "宿泊費（1人1泊あたり）"
    ws_cost["B6"] = data.lodging_fee
    ws_cost["B6"].fill = input_fill

    deposit_amount_row = 7
    ws_cost.cell(row=deposit_amount_row, column=1, value="前金（1人あたり）")
    ws_cost.cell(row=deposit_amount_row, column=2, value=data.deposit_amount).fill = input_fill

    breakfast_price_row = 8
    lunch_price_row = 9
    dinner_price_row = 10
    ws_cost.cell(row=breakfast_price_row, column=1, value="朝食単価")
    ws_cost.cell(row=breakfast_price_row, column=2, value=data.meal_prices.get("朝", 0)).fill = input_fill
    ws_cost.cell(row=lunch_price_row, column=1, value="昼食単価")
    ws_cost.cell(row=lunch_price_row, column=2, value=data.meal_prices.get("昼", 0)).fill = input_fill
    ws_cost.cell(row=dinner_price_row, column=1, value="夕食単価（通常時）")
    ws_cost.cell(row=dinner_price_row, column=2, value=data.meal_prices.get("夜", 0)).fill = input_fill

    # BBQ等、特定日だけ夕食単価が異なる場合の一覧（名簿シートの単価行から参照される）
    row = dinner_price_row + 2
    special_price_row_by_date: Dict[str, int] = {}
    if data.special_dinner_prices:
        ws_cost.cell(row=row, column=1, value="特別日（BBQ等）の夕食単価")
        ws_cost.cell(row=row, column=1).font = header_font
        ws_cost.cell(row=row, column=1).fill = header_fill
        ws_cost.cell(row=row, column=2, value="単価")
        ws_cost.cell(row=row, column=2).font = header_font
        ws_cost.cell(row=row, column=2).fill = header_fill
        row += 1
        for sp in data.special_dinner_prices:
            ws_cost.cell(row=row, column=1, value=f"{_camp_date_label(sp.date)}の夕食")
            cell = ws_cost.cell(row=row, column=2, value=sp.price)
            cell.fill = input_fill
            special_price_row_by_date[sp.date] = row
            row += 1
        row += 1

    ws_cost.cell(row=row, column=1, value="項目")
    ws_cost.cell(row=row, column=2, value="金額（全体）")
    ws_cost.cell(row=row, column=1).font = header_font
    ws_cost.cell(row=row, column=1).fill = header_fill
    ws_cost.cell(row=row, column=2).font = header_font
    ws_cost.cell(row=row, column=2).fill = header_fill

    row += 1
    items_start = row
    for item in data.cost_items:
        ws_cost.cell(row=row, column=1, value=item.label)
        cell = ws_cost.cell(row=row, column=2, value=item.amount)
        cell.fill = input_fill
        row += 1
    items_end = max(row - 1, items_start)

    total_row = row + 1
    ws_cost.cell(row=total_row, column=1, value="固定費合計")
    ws_cost.cell(row=total_row, column=2, value=f"=SUM(B{items_start}:B{items_end})").font = calc_font

    per_person_row = total_row + 1
    ws_cost.cell(row=per_person_row, column=1, value="一人あたり固定費")
    ws_cost.cell(row=per_person_row, column=2, value=f"=B{total_row}/B5").font = calc_font

    ws_cost.column_dimensions["A"].width = 28
    ws_cost.column_dimensions["B"].width = 16

    # ==================================================
    # シート2: 名簿
    # ==================================================
    ws = wb.create_sheet("名簿")

    col_grade = 1
    col_name = 2
    meal_start_col = 3
    n_dates = len(data.dates)
    col_nights = meal_start_col + n_dates * len(CAMP_MEALS)
    col_fee = col_nights + 1
    col_deposit = col_fee + 1
    col_balance = col_deposit + 1

    meal_price_row_by_meal = {"朝": breakfast_price_row, "昼": lunch_price_row, "夜": dinner_price_row}

    ws.cell(row=1, column=col_grade, value="学年")
    ws.cell(row=1, column=col_name, value="氏名")
    ws.cell(row=1, column=col_nights, value="泊数")
    ws.cell(row=1, column=col_fee, value="徴収金額")
    ws.cell(row=1, column=col_deposit, value="前金徴収")
    ws.cell(row=1, column=col_balance, value="残額（当日徴収）")
    ws.merge_cells(start_row=1, start_column=col_grade, end_row=3, end_column=col_grade)
    ws.merge_cells(start_row=1, start_column=col_name, end_row=3, end_column=col_name)
    ws.merge_cells(start_row=1, start_column=col_nights, end_row=3, end_column=col_nights)
    ws.merge_cells(start_row=1, start_column=col_fee, end_row=3, end_column=col_fee)
    ws.merge_cells(start_row=1, start_column=col_deposit, end_row=3, end_column=col_deposit)
    ws.merge_cells(start_row=1, start_column=col_balance, end_row=3, end_column=col_balance)

    for i, iso in enumerate(data.dates):
        base_col = meal_start_col + i * len(CAMP_MEALS)
        ws.cell(row=1, column=base_col, value=_camp_date_label(iso))
        ws.merge_cells(start_row=1, start_column=base_col, end_row=1, end_column=base_col + len(CAMP_MEALS) - 1)
        for j, meal in enumerate(CAMP_MEALS):
            ws.cell(row=2, column=base_col + j, value=meal)
            # 単価行: BBQ等の特別日はその夕食単価を、それ以外は通常単価(合宿全体費用シート)を参照する
            price_cell = ws.cell(row=3, column=base_col + j)
            if meal == "夜" and iso in special_price_row_by_date:
                price_cell.value = f"=合宿全体費用!$B${special_price_row_by_date[iso]}"
            else:
                price_cell.value = f"=合宿全体費用!$B${meal_price_row_by_meal[meal]}"
            price_cell.font = calc_font
            price_cell.alignment = center
            price_cell.border = border
            price_cell.number_format = "#,##0"

    for col in range(1, col_balance + 1):
        for r in (1, 2):
            cell = ws.cell(row=r, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border

    for col in (col_grade, col_name, col_nights, col_fee, col_deposit, col_balance):
        cell = ws.cell(row=3, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    for i, member in enumerate(data.members):
        r = roster_start_row + i
        ws.cell(row=r, column=col_grade, value=member.grade).border = border
        ws.cell(row=r, column=col_name, value=member.name).border = border

        dinner_coords = []
        member_att = data.attendance.get(member.name) or {}
        for d_i, iso in enumerate(data.dates):
            base_col = meal_start_col + d_i * len(CAMP_MEALS)
            date_att = member_att.get(iso) or {}
            for m_i, meal in enumerate(CAMP_MEALS):
                mark = "○" if date_att.get(meal) else "×"
                cell = ws.cell(row=r, column=base_col + m_i, value=mark)
                cell.alignment = center
                cell.border = border
                cell.fill = input_fill
                # 泊数(宿泊費の対象): 最終日は夜に宿泊しないため、最終日の夕食は
                # ○であっても泊数には数えない(食費としては引き続き計算される)
                if meal == "夜" and d_i < n_dates - 1:
                    dinner_coords.append(cell.coordinate)

        nights_formula = "=" + "+".join(f'COUNTIF({c},"○")' for c in dinner_coords) if dinner_coords else "=0"
        nights_cell = ws.cell(row=r, column=col_nights, value=nights_formula)
        nights_cell.font = calc_font
        nights_cell.alignment = center
        nights_cell.border = border

        # 食費: その人が○をつけた食事マスだけ、対応する単価行(3行目)を合計する
        if n_dates > 0:
            att_range = f"{get_column_letter(meal_start_col)}{r}:{get_column_letter(col_nights - 1)}{r}"
            price_range = f"{get_column_letter(meal_start_col)}$3:{get_column_letter(col_nights - 1)}$3"
            meal_cost_term = f'+SUMPRODUCT(({att_range}="○")*{price_range})'
        else:
            meal_cost_term = ""

        fee_formula = (
            f"=ROUNDUP({nights_cell.coordinate}*合宿全体費用!$B$6"
            f"{meal_cost_term}"
            f"+合宿全体費用!$B${per_person_row},-3)"
        )
        fee_cell = ws.cell(row=r, column=col_fee, value=fee_formula)
        fee_cell.font = calc_font
        fee_cell.alignment = center
        fee_cell.border = border
        fee_cell.number_format = "#,##0"

        # 前金徴収(○/×入力) と 残額(当日徴収額 = 徴収金額 - 前金を払っていれば前金額)
        deposit_mark = "○" if data.deposit_paid.get(member.name) else "×"
        deposit_cell = ws.cell(row=r, column=col_deposit, value=deposit_mark)
        deposit_cell.alignment = center
        deposit_cell.border = border
        deposit_cell.fill = input_fill

        balance_formula = (
            f'=ROUNDUP({fee_cell.coordinate}-IF({deposit_cell.coordinate}="○",'
            f"合宿全体費用!$B${deposit_amount_row},0),0)"
        )
        balance_cell = ws.cell(row=r, column=col_balance, value=balance_formula)
        balance_cell.font = calc_font
        balance_cell.alignment = center
        balance_cell.border = border
        balance_cell.number_format = "#,##0"

    ws.column_dimensions[get_column_letter(col_grade)].width = 8
    ws.column_dimensions[get_column_letter(col_name)].width = 14
    for col in range(meal_start_col, col_nights):
        ws.column_dimensions[get_column_letter(col)].width = 5
    ws.column_dimensions[get_column_letter(col_nights)].width = 7
    ws.column_dimensions[get_column_letter(col_fee)].width = 12
    ws.column_dimensions[get_column_letter(col_deposit)].width = 10
    ws.column_dimensions[get_column_letter(col_balance)].width = 14
    ws.freeze_panes = ws.cell(row=roster_start_row, column=meal_start_col).coordinate

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = "合宿参加者名簿.xlsx"
    headers = {
        "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"
    }
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


# ==========================================
# Google Sheets API - 回答の自動取得
# ==========================================
@app.post("/get-responses")
async def get_responses(config: SheetConfig):
    import requests

    headers = {"Authorization": f"Bearer {config.access_token}"}

    res = requests.get(
        f"https://sheets.googleapis.com/v4/spreadsheets/{config.spreadsheet_id}/values/A:Z",
        headers=headers
    )

    if res.status_code != 200:
        return {"error": f"データの取得に失敗しました: {res.text}"}

    data = res.json()
    rows = data.get("values", [])

    if len(rows) < 2:
        return {"error": "回答がまだありません。"}

    headers_row = rows[0]
    members = []

    for row in rows[1:]:
        entry = {}
        for i, header in enumerate(headers_row):
            entry[header] = row[i] if i < len(row) else ""

        name = entry.get("名前", "").strip()
        station = entry.get("最寄り駅", "").strip()
        role = entry.get("参加形態", "乗客").strip()
        capacity = entry.get("定員（運転手の方のみ・数字で入力）", "4").strip()
        want_with = entry.get("一緒になりたい人（カンマ区切り）", "").strip()
        awkward_with = entry.get("気まずい人（カンマ区切り）", "").strip()

        if not name or not station:
            continue

        members.append({
            "name": name,
            "station": station,
            "can_drive": role == "運転手",
            "capacity": int(capacity) if capacity.isdigit() else 4,
            "want_with": [w.strip() for w in want_with.split(",") if w.strip()],
            "awkward_with": [a.strip() for a in awkward_with.split(",") if a.strip()],
        })

    return {"members": members, "count": len(members)}
